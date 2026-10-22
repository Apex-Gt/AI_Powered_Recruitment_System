import pytest

from app.core.config import settings
from app.pipeline.errors import NeedsReviewError, PermanentDocumentError
from app.pipeline.fields import extract_contact_fields
from app.pipeline.text_extraction import extract_text, sniff_type

RESUME = """Jane Doe
jane.doe@example.com | +1 415 555 0132
linkedin.com/in/janedoe | github.com/janedoe

Senior Engineer at Acme, Jan 2019 - Present
Built Python and PostgreSQL services. Led a team of five engineers and shipped a payments platform
that processed millions of transactions a day. Wrote Spring Boot and Kafka integrations, owned on-call.
"""


def _pdf(tmp_path, text=RESUME, name="r.pdf"):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=10)
    path = tmp_path / name
    doc.save(str(path))
    doc.close()
    return path


def _docx(tmp_path, text=RESUME, name="r.docx"):
    import docx

    d = docx.Document()
    for line in text.splitlines():
        d.add_paragraph(line)
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Python, SQL"
    path = tmp_path / name
    d.save(str(path))
    return path


def test_pdf_text_is_extracted_and_hashed(tmp_path):
    result = extract_text(str(_pdf(tmp_path)))
    assert result.method in {"pymupdf", "pypdf"}
    assert "jane.doe@example.com" in result.text
    assert len(result.text_hash) == 64


def test_docx_text_includes_tables(tmp_path):
    result = extract_text(str(_docx(tmp_path)))
    assert result.method == "docx"
    assert "Spring Boot" in result.text
    assert "Skills | Python, SQL" in result.text


def test_txt_resume(tmp_path):
    path = tmp_path / "r.txt"
    path.write_text(RESUME, encoding="utf-8")
    assert extract_text(str(path)).method == "txt"


def test_real_type_is_taken_from_content_not_extension(tmp_path):
    path = _pdf(tmp_path, name="looks_like.docx")
    assert sniff_type(path) == ".pdf"
    assert extract_text(str(path)).method in {"pymupdf", "pypdf"}


@pytest.mark.parametrize(
    "content,code",
    [
        (b"", "EMPTY_FILE"),
        (b"\xd0\xcf\x11\xe0\x00\x00\x00\x00", "UNSUPPORTED_TYPE"),  # legacy .doc
        (b"\x89PNG\r\n\x1a\n\x00\x00", "UNSUPPORTED_TYPE"),
        (b"%PDF-1.7\nthis is not a real pdf", "CORRUPT_FILE"),
    ],
)
def test_permanent_failures_are_classified(tmp_path, content, code):
    path = tmp_path / "bad.bin"
    path.write_bytes(content)
    with pytest.raises(PermanentDocumentError) as exc:
        extract_text(str(path))
    assert exc.value.code == code
    assert not exc.value.retryable


def test_encrypted_pdf_is_permanent(tmp_path):
    import pymupdf

    src = _pdf(tmp_path)
    doc = pymupdf.open(str(src))
    out = tmp_path / "locked.pdf"
    doc.save(str(out), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner")
    doc.close()
    with pytest.raises(PermanentDocumentError) as exc:
        extract_text(str(out))
    assert exc.value.code == "ENCRYPTED_FILE"


def test_oversized_file_is_rejected(tmp_path, monkeypatch):
    monkeypatch.setattr(settings.pipeline, "max_file_mb", 0)
    path = tmp_path / "big.txt"
    path.write_text("x" * 100)
    with pytest.raises(PermanentDocumentError) as exc:
        extract_text(str(path))
    assert exc.value.code == "FILE_TOO_LARGE"


def test_blank_pdf_goes_to_review_not_failure(tmp_path, monkeypatch):
    import pymupdf

    monkeypatch.setattr("app.pipeline.text_extraction.tesseract_available", lambda: False)
    doc = pymupdf.open()
    doc.new_page()
    path = tmp_path / "blank.pdf"
    doc.save(str(path))
    doc.close()
    with pytest.raises(NeedsReviewError) as exc:
        extract_text(str(path))
    assert exc.value.needs_review
    assert exc.value.code in {"SCANNED_PDF", "NO_TEXT"}


def test_contact_fields_are_read_deterministically():
    fields = extract_contact_fields(RESUME, "jane_doe_resume.pdf")
    assert fields.email == "jane.doe@example.com"
    assert fields.name == "Jane Doe"
    assert fields.phone and "415" in fields.phone
    assert fields.linkedin_url == "https://linkedin.com/in/janedoe"
    assert fields.github_url == "https://github.com/janedoe"


def test_name_falls_back_to_file_name():
    fields = extract_contact_fields("email: a@b.co\n12345", "john_smith_cv.pdf")
    assert fields.name == "John Smith"
