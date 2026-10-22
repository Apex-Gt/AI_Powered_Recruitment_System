"""Text extraction for PDF, DOCX and plain-text resumes.

PyMuPDF reads PDFs (falling back to pypdf if it is not installed), python-docx reads DOCX, and Tesseract OCR
is used only when a PDF has too little embedded text to be a real text PDF. Every failure is classified as
permanent (corrupt, encrypted, unsupported, too large) or needs-review (readable file, no usable text), so
one bad file never fails a batch and is never retried forever.
"""

from __future__ import annotations

import hashlib
import re
import shutil
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from app.core.config import settings
from app.core.logging import get_logger
from app.pipeline.errors import NeedsReviewError, PermanentDocumentError

logger = get_logger(__name__)

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
_NUL = re.compile(r"[\x00\x0b\x0c]")
_BLANKS = re.compile(r"[ \t]+")
_BLANK_LINES = re.compile(r"\n{3,}")


@dataclass
class ExtractedText:
    text: str
    method: str  # pymupdf | pypdf | docx | txt | ocr
    pages: int = 0
    ocr_used: bool = False
    text_hash: str = ""
    warnings: list[str] = field(default_factory=list)


def clean_text(raw: str) -> str:
    text = _NUL.sub(" ", raw or "").replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(_BLANKS.sub(" ", line).strip() for line in text.split("\n"))
    return _BLANK_LINES.sub("\n\n", text).strip()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def sniff_type(path: Path) -> str:
    """Decide the real type from magic bytes, not the extension a client chose."""
    with path.open("rb") as fh:
        head = fh.read(8)
    if head.startswith(b"%PDF-"):
        return ".pdf"
    if head.startswith(b"PK\x03\x04"):
        return ".docx" if _is_docx(path) else "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return ".doc"  # legacy Word binary
    return ".txt" if _looks_like_text(path) else "unknown"


def _is_docx(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as zf:
            return "word/document.xml" in zf.namelist()
    except zipfile.BadZipFile:
        return False


def _looks_like_text(path: Path) -> bool:
    sample = path.read_bytes()[:4096]
    if b"\x00" in sample:
        return False
    try:
        sample.decode("utf-8")
        return True
    except UnicodeDecodeError:
        return False


def extract_text(file_path: str) -> ExtractedText:
    """Blocking extraction. Call through ``asyncio.to_thread`` from async code."""
    path = Path(file_path)
    if not path.is_file():
        raise PermanentDocumentError("FILE_NOT_FOUND", f"File not found: {path.name}")

    size = path.stat().st_size
    if size == 0:
        raise PermanentDocumentError("EMPTY_FILE", "The file is empty")
    if size > settings.pipeline.max_file_mb * 1024 * 1024:
        raise PermanentDocumentError("FILE_TOO_LARGE", f"The file is larger than {settings.pipeline.max_file_mb} MB")

    kind = sniff_type(path)
    if kind not in ALLOWED_EXTENSIONS:
        raise PermanentDocumentError(
            "UNSUPPORTED_TYPE", "Only PDF, DOCX and TXT resumes are supported" + (" (legacy .doc is not)" if kind == ".doc" else "")
        )

    if kind == ".pdf":
        result = _extract_pdf(path)
    elif kind == ".docx":
        result = _extract_docx(path)
    else:
        result = ExtractedText(clean_text(path.read_text(encoding="utf-8", errors="replace")), "txt", pages=1)

    result.text = clean_text(result.text)
    if not result.text:
        raise NeedsReviewError("NO_TEXT", "No readable text could be extracted from this file", {"method": result.method})
    result.text_hash = sha256_text(result.text)
    return result


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #


def _extract_pdf(path: Path) -> ExtractedText:
    try:
        import pymupdf  # type: ignore
    except ImportError:
        try:
            import fitz as pymupdf  # type: ignore  # older PyMuPDF module name
        except ImportError:
            return _extract_pdf_pypdf(path)

    try:
        doc = pymupdf.open(str(path))
    except Exception as e:  # noqa: BLE001 - PyMuPDF raises several unrelated types for corrupt files
        raise PermanentDocumentError("CORRUPT_FILE", f"The PDF could not be opened: {e}") from e

    with doc:
        if doc.needs_pass or doc.is_encrypted:
            raise PermanentDocumentError("ENCRYPTED_FILE", "The PDF is password protected")
        pages = doc.page_count
        if pages == 0:
            raise PermanentDocumentError("EMPTY_FILE", "The PDF has no pages")
        try:
            text = "\n\n".join(page.get_text("text") for page in doc)
        except Exception as e:  # noqa: BLE001
            raise PermanentDocumentError("CORRUPT_FILE", f"The PDF text could not be read: {e}") from e

        result = ExtractedText(text, "pymupdf", pages=pages)
        if len(clean_text(text)) < settings.pipeline.ocr_min_chars:
            return _ocr_fallback(doc, result)
        return result


def _extract_pdf_pypdf(path: Path) -> ExtractedText:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
    except (PdfReadError, ValueError, OSError) as e:
        raise PermanentDocumentError("CORRUPT_FILE", f"The PDF could not be opened: {e}") from e
    if reader.is_encrypted:
        raise PermanentDocumentError("ENCRYPTED_FILE", "The PDF is password protected")
    text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    result = ExtractedText(text, "pypdf", pages=len(reader.pages), warnings=["PyMuPDF not installed; used pypdf"])
    if len(clean_text(text)) < settings.pipeline.ocr_min_chars:
        raise NeedsReviewError(
            "SCANNED_PDF", "The PDF has little or no embedded text and OCR needs PyMuPDF to rasterise pages", {"pages": result.pages}
        )
    return result


def tesseract_available() -> bool:
    if not settings.pipeline.ocr_enabled:
        return False
    try:
        import pytesseract  # type: ignore  # noqa: F401
    except ImportError:
        return False
    return shutil.which("tesseract") is not None


def _ocr_fallback(doc, current: ExtractedText) -> ExtractedText:
    """Run OCR over the first few pages of a PDF that has (almost) no embedded text."""
    if not tesseract_available():
        # Keep whatever little text there is: the profile stage can still use it, but flag the file for review.
        if clean_text(current.text):
            current.warnings.append("Low text yield and OCR is unavailable; results may be incomplete")
            return current
        raise NeedsReviewError(
            "SCANNED_PDF",
            "The PDF appears to be a scan and OCR is not available (install Tesseract and pytesseract, or enable PIPELINE_OCR_ENABLED)",
            {"pages": current.pages},
        )

    import io

    import pytesseract  # type: ignore
    from PIL import Image  # type: ignore

    chunks: list[str] = []
    limit = min(doc.page_count, settings.pipeline.ocr_max_pages)
    for index in range(limit):
        pixmap = doc[index].get_pixmap(dpi=200)
        image = Image.open(io.BytesIO(pixmap.tobytes("png")))
        chunks.append(pytesseract.image_to_string(image, lang=settings.pipeline.ocr_language))
    text = "\n\n".join(chunks)
    warnings = list(current.warnings)
    if doc.page_count > limit:
        warnings.append(f"OCR limited to the first {limit} of {doc.page_count} pages")
    if len(clean_text(text)) <= len(clean_text(current.text)):
        return ExtractedText(current.text, current.method, pages=current.pages, warnings=warnings)
    return ExtractedText(text, "ocr", pages=current.pages, ocr_used=True, warnings=warnings + ["Text recovered with OCR"])


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #


def _extract_docx(path: Path) -> ExtractedText:
    try:
        import docx  # type: ignore
    except ImportError as e:
        raise PermanentDocumentError("UNSUPPORTED_TYPE", "DOCX support is not installed (python-docx)") from e

    try:
        document = docx.Document(str(path))
    except Exception as e:  # noqa: BLE001 - python-docx raises PackageNotFoundError, KeyError, zip errors...
        raise PermanentDocumentError("CORRUPT_FILE", f"The DOCX could not be opened: {e}") from e

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            seen: set[int] = set()
            cells = []
            for cell in row.cells:
                if id(cell._tc) in seen:  # merged cells repeat the same element
                    continue
                seen.add(id(cell._tc))
                cells.append(cell.text.strip())
            parts.append(" | ".join(c for c in cells if c))
    return ExtractedText("\n".join(parts), "docx", pages=1)


@lru_cache(maxsize=64)
def _cached(file_path: str, mtime_ns: int, size: int) -> ExtractedText:
    return extract_text(file_path)


def extract_text_cached(file_path: str) -> ExtractedText:
    """Same as :func:`extract_text` but memoised per (path, mtime, size): later stages re-read the same file."""
    stat = Path(file_path).stat() if Path(file_path).exists() else None
    if stat is None:
        return extract_text(file_path)
    return _cached(file_path, stat.st_mtime_ns, stat.st_size)
